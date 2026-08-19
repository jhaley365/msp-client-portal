"""
Export FPC AWS Instances + Syncro Asset Info → Excel
=====================================================
Pulls EC2 instances for customer FPC from DynamoDB, then looks up each
device in Syncro by hostname, and writes a combined spreadsheet.

Run from /opt/msp-portal/backend:
    python scripts/export_fpc_aws_syncro.py

Output: fpc_aws_syncro_export.xlsx in the current directory.

Requirements (already in .venv):
    boto3, requests, openpyxl
"""

from __future__ import annotations

import os
import sys
import requests
import boto3
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Config ────────────────────────────────────────────────────────────────────

CUSTOMER_ID   = "FPC"
REGION        = "us-east-1"
OUTPUT_FILE   = "fpc_aws_syncro_export.xlsx"

SYNCRO_API_KEY    = os.environ.get("SYNCRO_API_KEY", "")
SYNCRO_SUBDOMAIN  = os.environ.get("SYNCRO_SUBDOMAIN", "haley365")
SYNCRO_BASE       = f"https://{SYNCRO_SUBDOMAIN}.syncromsp.com/api/v1"

if not SYNCRO_API_KEY:
    # Try loading from sync.env
    env_file = "/etc/msp-portal/sync.env"
    if os.path.exists(env_file):
        with open(env_file) as f:
            for line in f:
                line = line.strip()
                if line.startswith("SYNCRO_API_KEY="):
                    SYNCRO_API_KEY = line.split("=", 1)[1]
                if line.startswith("SYNCRO_SUBDOMAIN="):
                    SYNCRO_SUBDOMAIN = line.split("=", 1)[1]
                    SYNCRO_BASE = f"https://{SYNCRO_SUBDOMAIN}.syncromsp.com/api/v1"

HEADERS = {
    "Authorization": SYNCRO_API_KEY,
    "Content-Type": "application/json",
    "Accept": "application/json",
}

# ── Helpers ───────────────────────────────────────────────────────────────────

def fetch_ec2_instances() -> list[dict]:
    """Pull all FPC EC2 instances from DynamoDB."""
    print("Fetching EC2 instances from DynamoDB...")
    ddb = boto3.resource("dynamodb", region_name=REGION)
    tbl = ddb.Table("EC2Instances")
    items: list[dict] = []
    kwargs: dict = {
        "IndexName": "customer_id-index",
        "KeyConditionExpression": boto3.dynamodb.conditions.Key("customer_id").eq(CUSTOMER_ID),
    }
    while True:
        resp = tbl.query(**kwargs)
        items.extend(resp.get("Items", []))
        if not resp.get("LastEvaluatedKey"):
            break
        kwargs["ExclusiveStartKey"] = resp["LastEvaluatedKey"]
    print(f"  Found {len(items)} EC2 instances")
    return sorted(items, key=lambda x: (x.get("name_tag") or x.get("instance_id", "")).lower())


def fetch_syncro_assets() -> list[dict]:
    """Fetch all assets from Syncro (all pages)."""
    print("Fetching assets from Syncro...")
    assets: list[dict] = []
    page = 1
    while True:
        resp = requests.get(
            f"{SYNCRO_BASE}/customer_assets",
            headers=HEADERS,
            params={"page": page, "per_page": 100},
            timeout=30,
        )
        if resp.status_code != 200:
            print(f"  Warning: Syncro assets returned {resp.status_code}: {resp.text[:200]}")
            break
        data = resp.json()
        batch = data.get("assets", [])
        if not batch:
            break
        assets.extend(batch)
        meta = data.get("meta", {})
        total_pages = meta.get("total_pages", 1)
        print(f"  Page {page}/{total_pages} — {len(assets)} assets so far")
        if page >= total_pages:
            break
        page += 1
    print(f"  Total Syncro assets: {len(assets)}")
    return assets


def build_asset_lookup(assets: list[dict]) -> dict[str, dict]:
    """Index Syncro assets by normalised hostname for fast lookup."""
    lookup: dict[str, dict] = {}
    for a in assets:
        name = (a.get("name") or a.get("hostname") or "").strip().upper()
        if name:
            lookup[name] = a
        # Also index by asset_tag if present
        tag = (a.get("asset_tag") or "").strip().upper()
        if tag and tag not in lookup:
            lookup[tag] = a
    return lookup


def match_asset(instance: dict, lookup: dict[str, dict]) -> dict | None:
    """Try to match an EC2 instance to a Syncro asset by name."""
    name = (instance.get("name_tag") or "").strip().upper()
    if name in lookup:
        return lookup[name]
    # Try without prefix (e.g. "FPC-DC-1.1" → "DC-1.1")
    parts = name.split("-", 1)
    if len(parts) == 2 and parts[1] in lookup:
        return lookup[parts[1]]
    return None


def extract_hardware(asset: dict) -> dict:
    """Pull hardware fields from a Syncro asset record."""
    if not asset:
        return {}

    # Syncro stores hardware info in the 'properties' dict or top-level fields
    props = asset.get("properties") or {}

    def get(*keys):
        for k in keys:
            v = asset.get(k) or props.get(k)
            if v:
                return str(v)
        return ""

    return {
        "syncro_name":      get("name", "hostname"),
        "syncro_id":        str(asset.get("id", "")),
        "os":               get("os_version", "operating_system", "os"),
        "os_version":       get("windows_version", "os_build"),
        "cpu":              get("cpu", "processor"),
        "cpu_cores":        get("cpu_count", "cores"),
        "ram_gb":           get("ram", "memory_gb", "memory"),
        "storage":          get("hdd_storage", "disk_size", "storage"),
        "last_seen":        get("last_seen_at", "updated_at"),
        "serial":           get("serial", "serial_number"),
        "manufacturer":     get("manufacturer"),
        "model":            get("model"),
        "syncro_customer":  get("customer_name"),
    }


# ── Excel builder ─────────────────────────────────────────────────────────────

DARK_BLUE  = "0C111E"
MID_BLUE   = "1E3A5F"
ACCENT     = "2F6BFF"
LIGHT_GREY = "F1F5F9"
WHITE      = "FFFFFF"
GREEN      = "16A34A"
ORANGE     = "D97706"

def col_border():
    thin = Side(style="thin", color="CBD5E1")
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def header_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def build_excel(instances: list[dict], asset_lookup: dict[str, dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "FPC AWS + Syncro"

    # ── Column definitions ────────────────────────────────────────────────────
    columns = [
        # (header, width, field_fn)
        ("Name",             24, lambda i, a: i.get("name_tag") or i.get("instance_id", "")),
        ("Instance ID",      22, lambda i, a: i.get("instance_id", "")),
        ("Type",             14, lambda i, a: i.get("instance_type", "")),
        ("State",            11, lambda i, a: i.get("state", "")),
        ("Region",           12, lambda i, a: i.get("region", "")),
        ("Private IP",       15, lambda i, a: i.get("private_ip", "")),
        ("Public IP",        15, lambda i, a: i.get("public_ip", "")),
        ("Platform",         10, lambda i, a: i.get("platform", "")),
        ("Launched",         13, lambda i, a: str(i.get("launch_time", ""))[:10]),
        # Divider — Syncro fields
        ("Syncro Name",      22, lambda i, a: a.get("syncro_name", "") if a else "Not Found"),
        ("OS",               30, lambda i, a: a.get("os", "") if a else ""),
        ("OS Version",       20, lambda i, a: a.get("os_version", "") if a else ""),
        ("CPU",              30, lambda i, a: a.get("cpu", "") if a else ""),
        ("CPU Cores",        11, lambda i, a: a.get("cpu_cores", "") if a else ""),
        ("RAM",              10, lambda i, a: a.get("ram_gb", "") if a else ""),
        ("Storage",          14, lambda i, a: a.get("storage", "") if a else ""),
        ("Manufacturer",     16, lambda i, a: a.get("manufacturer", "") if a else ""),
        ("Model",            20, lambda i, a: a.get("model", "") if a else ""),
        ("Serial",           18, lambda i, a: a.get("serial", "") if a else ""),
        ("Last Seen",        18, lambda i, a: a.get("last_seen", "") if a else ""),
    ]

    # ── Title row ─────────────────────────────────────────────────────────────
    ws.merge_cells(f"A1:{get_column_letter(len(columns))}1")
    title_cell = ws["A1"]
    title_cell.value = f"FPC — AWS Instances + Syncro Asset Info   |   Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    title_cell.font = Font(bold=True, color=WHITE, size=12)
    title_cell.fill = header_fill(DARK_BLUE)
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28

    # ── Section header row ────────────────────────────────────────────────────
    ws.merge_cells("A2:I2")
    aws_hdr = ws["A2"]
    aws_hdr.value = "AWS EC2"
    aws_hdr.font = Font(bold=True, color=WHITE, size=10)
    aws_hdr.fill = header_fill(ACCENT)
    aws_hdr.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"J2:{get_column_letter(len(columns))}2")
    syn_hdr = ws["J2"]
    syn_hdr.value = "Syncro Asset"
    syn_hdr.font = Font(bold=True, color=WHITE, size=10)
    syn_hdr.fill = header_fill(MID_BLUE)
    syn_hdr.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # ── Column headers ────────────────────────────────────────────────────────
    for col_idx, (header, width, _) in enumerate(columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        is_aws = col_idx <= 9
        cell.font = Font(bold=True, color=WHITE, size=9)
        cell.fill = header_fill("1E4080" if is_aws else "2D4A6E")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = col_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[3].height = 20

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_idx, instance in enumerate(instances, 4):
        asset_raw = match_asset(instance, asset_lookup)
        asset = extract_hardware(asset_raw) if asset_raw else None
        is_even = (row_idx % 2 == 0)

        for col_idx, (_, _, field_fn) in enumerate(columns, 1):
            value = field_fn(instance, asset)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(size=9)
            cell.border = col_border()
            cell.alignment = Alignment(vertical="center")

            # Row shading
            if is_even:
                cell.fill = PatternFill("solid", fgColor="EFF6FF")
            else:
                cell.fill = PatternFill("solid", fgColor=WHITE)

            # State colour
            if col_idx == 4:  # State column
                state = str(value).lower()
                if state == "running":
                    cell.font = Font(size=9, color=GREEN, bold=True)
                elif state in ("stopped", "terminated"):
                    cell.font = Font(size=9, color="DC2626", bold=True)

            # Not Found highlight
            if col_idx == 10 and value == "Not Found":
                cell.font = Font(size=9, color=ORANGE, bold=True)

        ws.row_dimensions[row_idx].height = 16

    # ── Freeze panes + auto-filter ────────────────────────────────────────────
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(columns))}{3 + len(instances)}"

    wb.save(OUTPUT_FILE)
    print(f"\nSaved: {OUTPUT_FILE}  ({len(instances)} rows)")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    if not SYNCRO_API_KEY:
        print("ERROR: SYNCRO_API_KEY not found in environment or /etc/msp-portal/sync.env")
        sys.exit(1)

    instances = fetch_ec2_instances()
    if not instances:
        print("No EC2 instances found for FPC — check DynamoDB EC2Instances table.")
        sys.exit(1)

    assets = fetch_syncro_assets()
    lookup = build_asset_lookup(assets)

    matched = sum(1 for i in instances if match_asset(i, lookup))
    print(f"\nMatched {matched}/{len(instances)} instances to Syncro assets")

    build_excel(instances, lookup)


if __name__ == "__main__":
    main()
